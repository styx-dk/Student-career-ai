import io
import mimetypes
import re
from pathlib import Path
from typing import BinaryIO

import fitz
from docx import Document as DocxDocument
from fastapi import HTTPException, status
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".ppt", ".pptx", ".txt", ".jpg", ".jpeg", ".png", ".webp"}
SUPPORTED_MIME_PREFIXES = ("application/pdf", "text/plain", "image/")
SUPPORTED_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "application/vnd.ms-powerpoint",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}


def safe_filename(filename: str) -> str:
    name = Path(filename).name
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", name).strip(".-")
    return stem[:180] or "document"


def validate_upload(filename: str, mime_type: str, size: int, max_bytes: int) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported file type. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}",
        )
    guessed = mimetypes.guess_type(filename)[0]
    mime_ok = mime_type in SUPPORTED_MIMES or mime_type.startswith(SUPPORTED_MIME_PREFIXES)
    if not mime_ok or (guessed and guessed != mime_type and mime_type != "application/octet-stream"):
        raise HTTPException(status_code=415, detail="File extension and MIME type do not match")
    if size <= 0 or size > max_bytes:
        raise HTTPException(status_code=413, detail=f"File must be between 1 byte and {max_bytes} bytes")
    return extension


def extract_text(data: bytes, extension: str) -> tuple[str, bool]:
    if extension == ".pdf":
        pdf = fitz.open(stream=data, filetype="pdf")
        text = "\n".join(page.get_text() for page in pdf)
        return text.strip(), len(text.strip()) < 80
    if extension == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip(), False
    if extension == ".pptx":
        deck = Presentation(io.BytesIO(data))
        text = "\n".join(
            shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text")
        )
        return text.strip(), False
    if extension == ".txt":
        return data.decode("utf-8", errors="replace").strip(), False
    if extension in IMAGE_EXTENSIONS:
        return "", True
    if extension in {".doc", ".ppt"}:
        raise ValueError("Legacy DOC/PPT extraction requires LibreOffice conversion configured on the server")
    raise ValueError("Unsupported extraction format")

